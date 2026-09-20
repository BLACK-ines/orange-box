from datetime import datetime, timedelta, time
import calendar
import random
import string

import json

import openpyxl
from reportlab.pdfgen import canvas

from reportlab.lib.pagesizes import letter

from django.shortcuts import render, redirect
from django.contrib import messages
from django.contrib.auth.decorators import login_required, user_passes_test
from django.http import HttpResponse, JsonResponse
from django.utils import timezone

from django.db.models import Q

from .models import (
    Employee, AttendanceRecord, ExcelFile, UploadConflict,
    Leaves, LeaveUsage, GracePeriod, MonthlyReport, MonthlyReportEntry,
    YearlyTrendReport, TrendCategory, Notification, InvitationCode, HRStaff,
    Department
)


def is_attendance_lead(user):
    return hasattr(user, 'attendancelead')


@login_required
def dashboard(request):
    return render(request, 'attendance/dashboard.html')



@login_required
def upload_excel(request):
    departments = Department.objects.all()

    if request.method == 'POST':
        uploaded_file = request.FILES.get('excel_file')
        department_id = request.POST.get('department')

        if uploaded_file and department_id:
            department = Department.objects.get(id=department_id)

            excel_record = ExcelFile.objects.create(
                file=uploaded_file,
                status='pending',
                uploaded_by=request.user
            )

            workbook = openpyxl.load_workbook(excel_record.file.path)
            sheet = workbook.active

            # ---------- PASS 1: VALIDATE EVERYTHING FIRST ----------
            parsed_rows = []
            for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
                if row[0] is None:
                    continue

                try:
                    emp_id, name, date, check_in, check_out = row
                except ValueError:
                    excel_record.delete()
                    messages.error(
                        request,
                        f"Upload rejected — row {row_number} is malformed (wrong number of columns). "
                        f"Please review and fix this row in the Excel file before re-uploading."
                    )
                    return redirect('upload_excel')

                try:
                    date_obj = date.date() if hasattr(date, 'date') else datetime.strptime(str(date), '%Y-%m-%d').date()
                    check_in_obj = check_in if (check_in is None or hasattr(check_in, 'hour')) else datetime.strptime(str(check_in), '%H:%M').time()
                    check_out_obj = check_out if (check_out is None or hasattr(check_out, 'hour')) else datetime.strptime(str(check_out), '%H:%M').time()
                except (ValueError, TypeError):
                    excel_record.delete()
                    messages.error(
                        request,
                        f"Upload rejected — row {row_number} has an invalid date or time value. "
                        f"Please review and fix this row in the Excel file before re-uploading."
                    )
                    return redirect('upload_excel')

                if not emp_id or not name:
                    excel_record.delete()
                    messages.error(
                        request,
                        f"Upload rejected — row {row_number} is missing an employee ID or name. "
                        f"Please review and fix this row in the Excel file before re-uploading."
                    )
                    return redirect('upload_excel')

                status = 'present' if check_in_obj else 'absent'
                total_hours = None
                if check_in_obj and check_out_obj:
                    duration = datetime.combine(date_obj, check_out_obj) - datetime.combine(date_obj, check_in_obj)
                    total_hours = duration.total_seconds() / 3600

                parsed_rows.append({
                    'emp_id': str(emp_id), 'name': name, 'date': date_obj,
                    'check_in': check_in_obj, 'check_out': check_out_obj,
                    'status': status, 'total_hours': total_hours
                })

            # ---------- PASS 2: EVERYTHING VALID — NOW ACTUALLY PROCESS ----------
            row_count = 0
            created_count = 0
            conflict_count = 0
            skipped_wrong_dept = []

            for r in parsed_rows:
                row_count += 1

                existing_employee = Employee.objects.filter(employee_id=r['emp_id']).first()
                if existing_employee and existing_employee.department_id != department.id:
                    skipped_wrong_dept.append(
                        f"{r['emp_id']} - {r['name']} (belongs to {existing_employee.department.name if existing_employee.department else 'no department'})"
                    )
                    continue

                employee, created = Employee.objects.get_or_create(
                    employee_id=r['emp_id'],
                    defaults={'name': r['name'], 'department': department}
                )

                existing_record = AttendanceRecord.objects.filter(employee=employee, date=r['date']).first()

                if existing_record:
                    if existing_record.matches(r['check_in'], r['check_out'], r['status']):
                        continue
                    else:
                        UploadConflict.objects.create(
                            employee=employee,
                            date=r['date'],
                            old_check_in=existing_record.check_in,
                            old_check_out=existing_record.check_out,
                            old_status=existing_record.status,
                            new_check_in=r['check_in'],
                            new_check_out=r['check_out'],
                            new_status=r['status'],
                            old_record=existing_record,
                            new_source_file=excel_record,
                            status='open'
                        )
                        Notification.objects.create(
                            type='conflict',
                            message=f"Conflict for {employee.name} ({employee.employee_id}) on {r['date']} — new file disagrees with existing record.",
                            status='unread',
                            recipient=request.user
                        )
                        conflict_count += 1
                else:
                    record = AttendanceRecord.objects.create(
                        employee=employee,
                        source_file=excel_record,
                        date=r['date'],
                        check_in=r['check_in'],
                        check_out=r['check_out'],
                        total_hours=r['total_hours'],
                        status=r['status'],
                        gap_resolved=(r['status'] != 'absent')
                    )
                    created_count += 1

                    if r['status'] == 'absent':
                        Notification.objects.create(
                            type='gap',
                            message=f"{employee.name} ({employee.employee_id}) was absent on {r['date']} — classify this gap.",
                            status='unread',
                            recipient=request.user,
                            related_record=record
                        )

            result_message = f"'{uploaded_file.name}' processed — {created_count} records created, {conflict_count} conflicts flagged out of {row_count} rows."
            if skipped_wrong_dept:
                result_message += f" Skipped {len(skipped_wrong_dept)} row(s) from a different department: {', '.join(skipped_wrong_dept)}."
                messages.warning(request, result_message)
            else:
                messages.success(request, result_message)

            return redirect('upload_excel')

    return render(request, 'attendance/upload.html', {'departments': departments})



@login_required
def conflict_list(request):
    conflicts = UploadConflict.objects.filter(status='open').select_related('employee', 'employee__department')
    return render(request, 'attendance/conflicts.html', {'conflicts': conflicts})


@user_passes_test(is_attendance_lead)
@login_required
def resolve_conflict(request, conflict_id):
    conflict = UploadConflict.objects.get(id=conflict_id)

    if request.method == 'POST':
        action = request.POST.get('action')

        if action == 'keep_old':
            conflict.resolution = 'kept_old'

        elif action == 'take_new':
            if conflict.old_record:
                conflict.old_record.check_in = conflict.new_check_in
                conflict.old_record.check_out = conflict.new_check_out
                conflict.old_record.status = conflict.new_status
                if conflict.new_check_in and conflict.new_check_out:
                    duration = datetime.combine(conflict.date, conflict.new_check_out) - datetime.combine(conflict.date, conflict.new_check_in)
                    conflict.old_record.total_hours = duration.total_seconds() / 3600
                else:
                    conflict.old_record.total_hours = None
                conflict.old_record.source_file = conflict.new_source_file
                conflict.old_record.save()
            conflict.resolution = 'took_new'

        elif action == 'delete_department_month':
            department = conflict.employee.department
            month = conflict.date.month
            year = conflict.date.year
            AttendanceRecord.objects.filter(
                employee__department=department, date__month=month, date__year=year
            ).delete()
            UploadConflict.objects.filter(
                employee__department=department, date__month=month, date__year=year
            ).update(status='resolved', resolution='department_month_deleted', resolved_date=datetime.now())
            messages.success(request, f"All {department.name} records for {month}/{year} deleted. Please re-upload a clean file.")
            return redirect('conflict_list')

        # Clean up related unread notifications for this employee
        Notification.objects.filter(
            type='conflict',
            message__icontains=conflict.employee.name,
            status='unread'
        ).update(status='resolved')

        conflict.status = 'resolved'
        conflict.resolved_date = datetime.now()
        conflict.resolved_by = request.user.attendancelead
        conflict.save()

        messages.success(request, f"Conflict for {conflict.employee.name} resolved ({conflict.resolution}).")
        return redirect('conflict_list')

    return render(request, 'attendance/resolve_conflict.html', {'conflict': conflict})


# ---------- LEAVES ----------

@login_required
def leave_list(request):
    leaves = Leaves.objects.all()
    return render(request, 'attendance/leave_list.html', {'leaves': leaves})


@login_required
def add_leave(request):
    if request.method == 'POST':
        Leaves.objects.create(
            name=request.POST.get('name'),
            description=request.POST.get('description'),
            category=request.POST.get('category'),
            max_hours_allowed=request.POST.get('max_hours_allowed'),
            limit_per_period=request.POST.get('limit_per_period'),
            period_type=request.POST.get('period_type'),
            is_paid=request.POST.get('is_paid') == 'on'
        )
        messages.success(request, "Leave type created.")
        return redirect('leave_list')

    return render(request, 'attendance/add_leave.html')


# ---------- LEAVE USAGE ----------

@login_required
def log_leave_usage(request):
    if request.method == 'POST':
        employee = Employee.objects.get(id=request.POST.get('employee'))
        leave_type = Leaves.objects.get(id=request.POST.get('leave_type'))
        hours_requested = float(request.POST.get('hours_used'))
        entry_date = request.POST.get('date')

        if hours_requested > leave_type.max_hours_allowed:
            messages.error(
                request,
                f"Cannot log {hours_requested}h — max allowed per entry for "
                f"'{leave_type.name}' is {leave_type.max_hours_allowed}h."
            )
            return redirect('log_leave_usage')

        entry_date_obj = datetime.strptime(entry_date, '%Y-%m-%d').date()

        if leave_type.period_type == 'month':
            existing_count = LeaveUsage.objects.filter(
                employee=employee, leave_type=leave_type,
                date__year=entry_date_obj.year, date__month=entry_date_obj.month
            ).count()
            period_label = "this month"
        else:
            existing_count = LeaveUsage.objects.filter(
                employee=employee, leave_type=leave_type,
                date__year=entry_date_obj.year
            ).count()
            period_label = "this year"

        if existing_count >= leave_type.limit_per_period:
            messages.error(
                request,
                f"'{employee.name}' has already reached the limit of "
                f"{leave_type.limit_per_period} uses for '{leave_type.name}' {period_label}."
            )
            return redirect('log_leave_usage')

        LeaveUsage.objects.create(
            employee=employee,
            leave_type=leave_type,
            date=entry_date_obj,
            hours_used=hours_requested
        )
        messages.success(request, f"Leave logged for {employee.name}.")
        return redirect('log_leave_usage')

    employees = Employee.objects.all()
    leaves = Leaves.objects.all()
    usages = LeaveUsage.objects.all().order_by('-date')

    today = datetime.now()
    balances = []
    for emp in employees:
        for leave in leaves:
            if leave.period_type == 'month':
                used_count = LeaveUsage.objects.filter(
                    employee=emp, leave_type=leave, date__year=today.year, date__month=today.month
                ).count()
            else:
                used_count = LeaveUsage.objects.filter(
                    employee=emp, leave_type=leave, date__year=today.year
                ).count()
            remaining = leave.limit_per_period - used_count
            balances.append({
                'employee': emp.name,
                'leave_type': leave.name,
                'used': used_count,
                'limit': leave.limit_per_period,
                'period': leave.get_period_type_display(),
                'remaining': remaining
            })

    return render(request, 'attendance/log_leave_usage.html', {
        'employees': employees,
        'leaves': leaves,
        'usages': usages,
        'balances': balances
    })






# ---------- GRACE PERIOD ----------
@user_passes_test(is_attendance_lead)
@login_required
def grace_period_list(request):
    departments = Department.objects.all()
    periods = GracePeriod.objects.all()

    if request.method == 'POST':
        department = Department.objects.get(id=request.POST.get('department'))
        GracePeriod.objects.update_or_create(
            department=department,
            defaults={
                'minutes': request.POST.get('minutes'),
                'start_time': request.POST.get('start_time'),
                'end_time': request.POST.get('end_time'),
                'set_by': request.user.attendancelead if hasattr(request.user, 'attendancelead') else None
            }
        )
        messages.success(request, f"Time settings for {department.name} updated.")
        return redirect('grace_period_list')

    return render(request, 'attendance/grace_period_list.html', {'periods': periods, 'departments': departments})


# ---------- MONTHLY REPORT ----------

@login_required
def generate_monthly_report(request):
    departments = Department.objects.all()
    search = request.GET.get('search', '')

    if request.method == 'POST':
        department_id = request.POST.get('department')
        month = int(request.POST.get('month'))
        year = int(request.POST.get('year'))
        department = Department.objects.get(id=department_id)

        # block if this department+month already has a report
        if MonthlyReport.objects.filter(department=department, month=month, year=year).exists():
            messages.error(request, f"A report for {department.name} - {month}/{year} already exists. View it instead.")
            return redirect('generate_monthly_report')

        # block if unresolved conflicts exist for this department
        open_conflicts = UploadConflict.objects.filter(employee__department=department, status='open').exists()
        if open_conflicts:
            messages.error(request, f"Cannot generate report — {department.name} has unresolved conflicts. Resolve them first.")
            return redirect('conflict_list')

        report = MonthlyReport.objects.create(
            department=department, month=month, year=year, generated_by=request.user
        )

        days_in_month = calendar.monthrange(year, month)[1]
        employees = Employee.objects.filter(department=department)

        service_leave = Leaves.objects.filter(name__iexact='service').first()
        mission_leave = Leaves.objects.filter(name__iexact='mission').first()

        for employee in employees:
            records = AttendanceRecord.objects.filter(employee=employee, date__month=month, date__year=year)
            total_worked_hours = sum(r.total_hours or 0 for r in records)
            days_worked = records.filter(status='present').count()
            days_absent = records.filter(status='absent').count()

            grace_period = GracePeriod.objects.filter(department=department).order_by('-updated_date').first()
            grace_minutes = grace_period.minutes if grace_period else 0
            standard_start = grace_period.start_time if grace_period else time(8, 0)

            late_threshold_minutes = (standard_start.hour * 60 + standard_start.minute) + grace_minutes
            late_threshold = time(late_threshold_minutes // 60, late_threshold_minutes % 60)
            lateness_count = records.filter(status='present', check_in__gt=late_threshold).count()
            late_records = records.filter(status='present', check_in__gt=late_threshold)
            total_late_minutes = 0
            for r in late_records:
                diff = (r.check_in.hour * 60 + r.check_in.minute) - (late_threshold.hour * 60 + late_threshold.minute)
                total_late_minutes += diff

            service_hours = 0
            mission_hours = 0
            if service_leave:
                service_hours = sum(u.hours_used for u in LeaveUsage.objects.filter(employee=employee, leave_type=service_leave, date__month=month, date__year=year))
            if mission_leave:
                mission_hours = sum(u.hours_used for u in LeaveUsage.objects.filter(employee=employee, leave_type=mission_leave, date__month=month, date__year=year))

            attendance_percentage = (days_worked / days_in_month) * 100 if days_in_month else 0

            # detect mid-month exception: first record for this employee starts after day 5 of the month
            first_record = AttendanceRecord.objects.filter(employee=employee).order_by('date').first()
            is_exception = bool(first_record and first_record.date.day > 5 and first_record.date.month == month and first_record.date.year == year)

            MonthlyReportEntry.objects.create(
                report=report,
                employee=employee,
                total_working_hours=total_worked_hours,
                total_lateness=lateness_count,
                total_late_minutes=total_late_minutes,
                total_absence=days_absent,
                total_working_days=days_worked,
                attendance_percentage=attendance_percentage,
                service_hours=service_hours,
                mission_hours=mission_hours,
                total_hours_with_leave=total_worked_hours + service_hours + mission_hours,
                is_exception=is_exception
            )

        messages.success(request, f"Report generated: {department.name} - {month}/{year}.")
        return redirect('monthly_report_detail', report_id=report.id)

    reports = MonthlyReport.objects.all().order_by('-year', '-month')
    if search:
        reports = reports.filter(department__name__icontains=search)

    return render(request, 'attendance/generate_monthly_report.html', {
        'reports': reports, 'departments': departments, 'search': search
    })


@login_required
def monthly_report_detail(request, report_id):
    report = MonthlyReport.objects.get(id=report_id)
    entries = report.entries.filter(is_exception=False)
    exceptions = report.entries.filter(is_exception=True)
    has_open_conflicts = UploadConflict.objects.filter(employee__department=report.department, status='open').exists()
    return render(request, 'attendance/monthly_report_detail.html', {
        'report': report, 'entries': entries, 'exceptions': exceptions, 'has_open_conflicts': has_open_conflicts
    })




@login_required
def export_monthly_report_excel(request, report_id):
    report = MonthlyReport.objects.get(id=report_id)
    entries = report.entries.all()

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Monthly Report"
    sheet.append([f"{report.department.name} - {report.month}/{report.year}"])
    sheet.append(['ID', 'Name', 'Service (h)', 'Mission (h)', 'Absence', 'Lateness', 'Total Late (min)', 'Worked (h)', 'Attendance %', 'Total (h)'])

    for entry in entries:
        sheet.append([
            entry.employee.employee_id,
            entry.employee.name,
            entry.service_hours,
            entry.mission_hours,
            entry.total_absence,
            entry.total_lateness,
            entry.total_late_minutes,
            entry.total_working_hours,
            round(entry.attendance_percentage, 1),
            entry.total_hours_with_leave
        ])

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{report.department.name}_{report.month}_{report.year}.xlsx"'
    workbook.save(response)
    return response




@login_required
def export_monthly_report_pdf(request, report_id):
    report = MonthlyReport.objects.get(id=report_id)
    entries = report.entries.filter(is_exception=False)
    exceptions = report.entries.filter(is_exception=True)

    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="{report.department.name}_{report.month}_{report.year}.pdf"'

    p = canvas.Canvas(response, pagesize=letter)
    width, height = letter

    p.setFillColorRGB(0.18, 0.32, 0.2)
    p.rect(0, height - 70, width, 70, fill=1, stroke=0)
    p.setFillColorRGB(1, 1, 1)
    p.setFont("Helvetica-Bold", 16)
    p.drawString(40, height - 35, f"{report.department.name} — Monthly Report")
    p.setFont("Helvetica", 10)
    p.drawString(40, height - 52, f"{report.month}/{report.year}  |  Generated {report.generated_date.strftime('%d %b %Y')}")

    y = height - 100
    p.setFillColorRGB(0, 0, 0)

    def draw_table(title, data, y):
        p.setFont("Helvetica-Bold", 12)
        p.drawString(40, y, title)
        y -= 18
        headers = ['ID', 'Name', 'Serv.', 'Mission', 'Absence', 'Late', 'Late Min', 'Worked', 'Attend %', 'Total']
        col_x = [40, 85, 195, 230, 275, 320, 355, 400, 445, 500]
        p.setFont("Helvetica-Bold", 8)
        p.setFillColorRGB(0.9, 0.95, 0.9)
        p.rect(38, y - 4, width - 78, 14, fill=1, stroke=0)
        p.setFillColorRGB(0, 0, 0)
        for i, h in enumerate(headers):
            p.drawString(col_x[i], y, h)
        y -= 16
        p.setFont("Helvetica", 8)
        for entry in data:
            if y < 60:
                p.showPage()
                y = height - 60
            row = [entry.employee.employee_id, entry.employee.name[:16], str(entry.service_hours),
                   str(entry.mission_hours), str(entry.total_absence), str(entry.total_lateness),
                   str(entry.total_late_minutes), str(entry.total_working_hours),
                   f"{round(entry.attendance_percentage,1)}%", str(entry.total_hours_with_leave)]
            for i, val in enumerate(row):
                p.drawString(col_x[i], y, val)
            y -= 14
        return y

    y = draw_table("Employees", entries, y)

    if exceptions:
        y -= 15
        y = draw_table("Joined Mid-Month", exceptions, y)

    p.showPage()
    p.save()
    return response






# ---------- TREND CATEGORIES ----------

@login_required
def trend_category_list(request):
    if request.method == 'POST':
        TrendCategory.objects.create(
            name=request.POST.get('name'),
            description=request.POST.get('description')
        )
        messages.success(request, "Trend category created.")
        return redirect('trend_category_list')

    categories = TrendCategory.objects.all()
    return render(request, 'attendance/trend_category_list.html', {'categories': categories})


# ---------- YEARLY TREND REPORT ----------
@user_passes_test(is_attendance_lead)
@login_required
def generate_yearly_report(request):
    departments = Department.objects.all()
    leave_types = Leaves.objects.all()

    if request.method == 'POST':
        department_id = request.POST.get('department')
        year = int(request.POST.get('year'))
        output_type = request.POST.get('output_type')
        category_ids = request.POST.getlist('categories')
        department = Department.objects.get(id=department_id)

        existing_count = YearlyTrendReport.objects.filter(department=department, year=year, output_type=output_type).count()
        report = YearlyTrendReport.objects.create(
            department=department, year=year, output_type=output_type,
            generated_by=request.user.attendancelead, display_number=existing_count + 1
        )
        report.categories.set(category_ids)

        messages.success(request, f"Yearly report generated: {department.name} - {year} ({output_type}).")
        return redirect('yearly_report_detail', report_id=report.id)

    reports = YearlyTrendReport.objects.all().order_by('-year')
    return render(request, 'attendance/generate_yearly_report.html', {
        'reports': reports, 'departments': departments, 'leave_types': leave_types
    })





@login_required
def yearly_report_detail(request, report_id):
    report = YearlyTrendReport.objects.get(id=report_id)

    # month-by-month attendance/hours summary
    monthly_summaries = []
    for month in range(1, 13):
        entries = MonthlyReportEntry.objects.filter(
            report__department=report.department, report__year=report.year, report__month=month
        )
        if entries.exists():
            avg_attendance = sum(e.attendance_percentage for e in entries) / entries.count()
            total_hours = sum(e.total_working_hours for e in entries)
            monthly_summaries.append({'month': month, 'avg_attendance': round(avg_attendance, 1), 'total_hours': total_hours})

    # leave usage by category by month (for the bar/line chart)
    categories = report.categories.all()
    leave_chart_data = {}
    for category in categories:
        monthly_counts = []
        for month in range(1, 13):
            count = LeaveUsage.objects.filter(
                leave_type=category, date__year=report.year, date__month=month,
                employee__department=report.department
            ).values('employee').distinct().count()
            monthly_counts.append(count)
        leave_chart_data[category.name] = monthly_counts

    # on-time vs late pie chart data
    all_year_entries = MonthlyReportEntry.objects.filter(report__department=report.department, report__year=report.year)
    total_late = sum(e.total_lateness for e in all_year_entries)
    total_working_days = sum(e.total_working_days for e in all_year_entries)
    total_on_time = max(total_working_days - total_late, 0)

    return render(request, 'attendance/yearly_report_detail.html', {
        'report': report,
        'summaries': monthly_summaries,
        'leave_chart_data': json.dumps(leave_chart_data),
        'total_late': total_late,
        'total_on_time': total_on_time,
    })




# ---------- NOTIFICATIONS ----------

@login_required
def notification_list(request):
    notifications = Notification.objects.filter(recipient=request.user).order_by('-date')
    return render(request, 'attendance/notifications.html', {'notifications': notifications})


@login_required
def mark_notification_read(request, notification_id):
    notification = Notification.objects.get(id=notification_id)
    notification.status = 'resolved'
    notification.save()
    return redirect('notification_list')


# ---------- INVITATION CODE ----------
@user_passes_test(is_attendance_lead)
@login_required
def generate_invite(request):
    if request.method == 'POST':
        code = ''.join(random.choices(string.ascii_uppercase + string.digits, k=20))
        InvitationCode.objects.create(
            code=code,
            expiry_date=timezone.now() + timedelta(minutes=30),
            status='active',
            generated_by=request.user.attendancelead
        )
        messages.success(request, f"Invitation code generated: {code} (expires in 30 min)")
        return redirect('generate_invite')

    codes = InvitationCode.objects.all().order_by('-generated_date')
    return render(request, 'attendance/generate_invite.html', {'codes': codes})


def signup(request):
    if request.method == 'POST':
        code_str = request.POST.get('code')
        name = request.POST.get('full_name')
        email = request.POST.get('email')
        password = request.POST.get('password')

        try:
            invite = InvitationCode.objects.get(code=code_str, status='active')

            if invite.expiry_date < timezone.now():
                messages.error(request, "This code has expired.")
                return redirect('signup')

            HRStaff.objects.create_user(email=email, password=password, full_name=name)
            invite.status = 'used'
            invite.save()

            messages.success(request, "Account created! You can now log in.")
            return redirect('login')

        except InvitationCode.DoesNotExist:
            messages.error(request, "Invalid invitation code.")
            return redirect('signup')

    return render(request, 'attendance/signup.html')


@user_passes_test(is_attendance_lead)
@login_required
def team_list(request):
    staff = HRStaff.objects.all().order_by('date_joined')
    return render(request, 'attendance/team_list.html', {'staff': staff})


@user_passes_test(is_attendance_lead)
@login_required
def generate_invite_ajax(request):
    code = ''.join(random.choices(string.ascii_uppercase + string.digits, k=20))
    InvitationCode.objects.create(
        code=code,
        expiry_date=timezone.now() + timedelta(minutes=30),
        status='active',
        generated_by=request.user.attendancelead
    )
    return JsonResponse({'code': code})






# ---------- EMPLOYEES ----------

@login_required
def employee_search(request):
    query = request.GET.get('q', '')
    results = []
    if query:
        results = Employee.objects.filter(
            Q(name__icontains=query) | Q(employee_id__icontains=query)
        )
    return render(request, 'attendance/employee_search.html', {'results': results, 'query': query})


@login_required
def employee_detail(request, employee_id):
    employee = Employee.objects.get(id=employee_id)

    if request.method == 'POST':
        employee.name = request.POST.get('name')
        employee.sex = request.POST.get('sex')
        employee.birthday = request.POST.get('birthday') or None
        employee.email = request.POST.get('email') or None
        employee.phone_number = request.POST.get('phone_number')
        dept_id = request.POST.get('department')
        employee.department = Department.objects.get(id=dept_id) if dept_id else None
        employee.save()
        messages.success(request, "Employee details updated.")
        return redirect('employee_detail', employee_id=employee.id)

    entries = MonthlyReportEntry.objects.filter(employee=employee).select_related('report').order_by('report__year', 'report__month')
    departments = Department.objects.all()

    return render(request, 'attendance/employee_detail.html', {
        'employee': employee, 'entries': entries, 'departments': departments
    })


@login_required
def export_employee_history_pdf(request, employee_id):
    employee = Employee.objects.get(id=employee_id)
    entries = MonthlyReportEntry.objects.filter(employee=employee).select_related('report').order_by('report__year', 'report__month')

    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="{employee.employee_id}_history.pdf"'

    p = canvas.Canvas(response, pagesize=letter)
    width, height = letter

    # header bar
    p.setFillColorRGB(0.18, 0.32, 0.2)
    p.rect(0, height - 70, width, 70, fill=1, stroke=0)
    p.setFillColorRGB(1, 1, 1)
    p.setFont("Helvetica-Bold", 16)
    p.drawString(40, height - 35, f"{employee.name} ({employee.employee_id})")
    p.setFont("Helvetica", 10)
    dept_name = employee.department.name if employee.department else "-"
    p.drawString(40, height - 52, f"Department: {dept_name}  |  Employee History Report")

    y = height - 100
    p.setFillColorRGB(0, 0, 0)

    p.setFont("Helvetica-Bold", 12)
    p.drawString(40, y, "Monthly History")
    y -= 18

    headers = ['Month', 'Year', 'Worked (h)', 'Absence', 'Lateness', 'Attendance %']
    col_x = [40, 110, 160, 240, 310, 380]
    p.setFont("Helvetica-Bold", 8)
    p.setFillColorRGB(0.9, 0.95, 0.9)
    p.rect(38, y - 4, width - 78, 14, fill=1, stroke=0)
    p.setFillColorRGB(0, 0, 0)
    for i, h in enumerate(headers):
        p.drawString(col_x[i], y, h)
    y -= 16

    p.setFont("Helvetica", 8)
    if entries:
        for entry in entries:
            if y < 60:
                p.showPage()
                y = height - 60
            row = [
                str(entry.report.month), str(entry.report.year),
                str(entry.total_working_hours), str(entry.total_absence),
                str(entry.total_lateness), f"{round(entry.attendance_percentage, 1)}%"
            ]
            for i, val in enumerate(row):
                p.drawString(col_x[i], y, val)
            y -= 14
    else:
        p.drawString(40, y, "No report history available for this employee yet.")

    p.showPage()
    p.save()
    return response





@user_passes_test(is_attendance_lead)
@login_required
def delete_monthly_report(request, report_id):
    report = MonthlyReport.objects.get(id=report_id)
    if request.method == 'POST':
        dept_name = report.department.name
        month, year = report.month, report.year
        report.delete()
        messages.success(request, f"Report {dept_name} - {month}/{year} deleted.")
        return redirect('generate_monthly_report')
    return redirect('monthly_report_detail', report_id=report.id)





@user_passes_test(is_attendance_lead)
@login_required
def delete_yearly_report(request, report_id):
    report = YearlyTrendReport.objects.get(id=report_id)
    if request.method == 'POST':
        dept_name = report.department.name
        year = report.year
        report.delete()
        messages.success(request, f"Yearly report {dept_name} - {year} deleted.")
        return redirect('generate_yearly_report')
    return redirect('yearly_report_detail', report_id=report.id)



@login_required
def resolve_gap(request, record_id):
    record = AttendanceRecord.objects.get(id=record_id)
    leave_types = Leaves.objects.all()

    if request.method == 'POST':
        choice = request.POST.get('choice')

        if choice == 'leave':
            leave_type_id = request.POST.get('leave_type')
            new_leave_name = request.POST.get('new_leave_name')

            if new_leave_name:
                leave_type = Leaves.objects.create(
                    name=new_leave_name, category='general',
                    max_hours_allowed=8, limit_per_period=999, is_paid=False
                )
            else:
                leave_type = Leaves.objects.get(id=leave_type_id)

            LeaveUsage.objects.create(
                employee=record.employee, leave_type=leave_type,
                date=record.date, hours_used=8
            )

        record.gap_resolved = True
        record.save()

        Notification.objects.filter(
            type='gap', message__icontains=record.employee.name, status='unread'
        ).update(status='resolved')

        messages.success(request, "Gap resolved.")
        return redirect('notification_list')

    return render(request, 'attendance/resolve_gap.html', {'record': record, 'leave_types': leave_types})


@login_required
def create_department_ajax(request):
    if request.method == 'POST':
        name = request.POST.get('name')
        start_time = request.POST.get('start_time') or '08:00'
        end_time = request.POST.get('end_time') or '17:00'
        minutes = request.POST.get('minutes') or 0

        if name:
            dept, created = Department.objects.get_or_create(name=name)
            if created:
                GracePeriod.objects.create(
                    department=dept, start_time=start_time, end_time=end_time, minutes=minutes,
                    set_by=request.user.attendancelead if hasattr(request.user, 'attendancelead') else None
                )
            return JsonResponse({'id': dept.id, 'name': dept.name})
    return JsonResponse({'error': 'invalid'}, status=400)



